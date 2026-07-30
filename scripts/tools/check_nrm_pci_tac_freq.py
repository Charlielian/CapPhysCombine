#!/usr/bin/env python3
"""Probe NRM Excel for PCI/TAC/freq and match against 原始小区表."""
from __future__ import annotations

import re
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook

_PROJECT_ROOT = Path(__file__).resolve().parents[2]
BASE = _PROJECT_ROOT / "data" / "网管配置"
DB = _PROJECT_ROOT / "capphys_unified.db"


def read_cols(path: Path, sheet: str, cols: list[str]) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    first = next(it, None)
    idx = {str(h): i for i, h in enumerate(header) if h is not None}
    want = [c for c in cols if c in idx]
    out = {c: [] for c in want}

    def handle(row):
        if row is None:
            return
        for c in want:
            out[c].append(row[idx[c]] if idx[c] < len(row) else None)

    if first is not None:
        joined = " ".join(str(x) for x in first if x is not None)
        if not re.search(r"[\u4e00-\u9fff]", joined):
            handle(first)
    for row in it:
        handle(row)
    wb.close()
    return pd.DataFrame(out)


def clean_num(s):
    if s is None or (isinstance(s, float) and pd.isna(s)):
        return None
    s = str(s).strip()
    if s == "" or s.lower() in ("nan", "none", "-"):
        return None
    if s.lower().startswith("0x"):
        try:
            return int(s, 16)
        except Exception:
            return s
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except Exception:
        return s


def plmn_norm(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    s = str(x).strip().replace("-", "")
    m = re.match(r"(\d{3})[-_]?(\d{2})", s)
    if m:
        return m.group(1) + m.group(2)
    digits = re.sub(r"\D", "", s)
    return digits or None


def make_cgi(plmn, node, cell):
    plmn = plmn_norm(plmn) or "46000"
    try:
        node = int(float(str(node)))
        cell = int(float(str(cell)))
    except Exception:
        return None
    return f"{plmn[:3]}-{plmn[3:]}-{node}-{cell}"


def variants(cgi):
    if cgi is None or (isinstance(cgi, float) and pd.isna(cgi)):
        return set()
    s = str(cgi).strip().upper().replace(" ", "").replace("_", "-")
    out = {s, s.replace("-", "")}
    parts = re.split(r"-", s)
    if len(parts) == 4:
        mcc, mnc, node, cell = parts
        out.add(f"{mcc}-{mnc}-{node}-{cell}")
        out.add(f"{mcc}{mnc}-{node}-{cell}")
        if mnc.isdigit():
            out.add(f"{mcc}-{int(mnc)}-{node}-{cell}")
            out.add(f"{mcc}-{int(mnc):02d}-{node}-{cell}")
        if node.isdigit() and cell.isdigit():
            out.add(f"{int(node)}-{int(cell)}")
    return out


def main():
    con = duckdb.connect(str(DB), read_only=True)
    phys = con.execute('SELECT CGI, 网络制式, 小区名称, BAND, BAND_A, 频点 FROM "原始小区表"').df()
    con.close()
    print("现网", len(phys))
    print(phys["网络制式"].value_counts(dropna=False).head(15).to_string())
    print("CGI样例", phys["CGI"].head(8).tolist())

    nr_parts = []
    for fname in ["2.6G.xlsx", "700M.xlsx"]:
        p = BASE / fname
        print(f"\n解析 {fname} ...", flush=True)
        ssb = read_cols(p, "CellDefiningSSB", ["ManagedElement", "ldn", "pci", "ssbFrequency"])
        print("  SSB", len(ssb), flush=True)
        du = read_cols(p, "NRCellDU", ["ManagedElement", "ldn", "moId", "cellLocalId", "nRTAC", "nrCarrierGroupId"])
        print("  DU", len(du), flush=True)
        cu = read_cols(p, "NRCellCU", ["ManagedElement", "ldn", "moId", "cellLocalId", "frequency", "ssbFrequency", "nrCarrierGroupId"])
        print("  CU", len(cu), flush=True)
        gnb = read_cols(p, "GNBCUCPFunction", ["ManagedElement", "moId", "gNBId", "gNBIdLength", "pLMNId"])
        print("  GNB", len(gnb), flush=True)
        phy = read_cols(p, "NRPhysicalCellDU", ["ManagedElement", "moId", "nrCarrierGroupId"])
        print("  PHY", len(phy), flush=True)

        def parent_phys(ldn):
            if ldn is None:
                return None
            m = re.search(r"NRPhysicalCellDU=([^,]+)", str(ldn))
            return m.group(1) if m else None

        ssb = ssb.copy()
        ssb["physMo"] = ssb["ldn"].map(parent_phys)
        ssb_g = ssb.groupby(["ManagedElement", "physMo"], as_index=False).agg({"pci": "first", "ssbFrequency": "first"})
        phy = phy.copy()
        phy["physMo"] = phy["moId"].astype(str)
        phy2 = phy.merge(ssb_g, on=["ManagedElement", "physMo"], how="left")
        gnb_map = gnb.groupby("ManagedElement", as_index=False).agg({"gNBId": "first", "gNBIdLength": "first", "pLMNId": "first"})
        du2 = du.merge(gnb_map, on="ManagedElement", how="left")
        merged = du2.merge(
            phy2[["ManagedElement", "nrCarrierGroupId", "pci", "ssbFrequency"]],
            on=["ManagedElement", "nrCarrierGroupId"],
            how="left",
        )
        if len(cu):
            merged = merged.merge(
                cu[["ManagedElement", "cellLocalId", "frequency", "ssbFrequency"]].rename(columns={"ssbFrequency": "ssb_cu"}),
                on=["ManagedElement", "cellLocalId"],
                how="left",
            )
            merged["ssbFrequency"] = merged["ssbFrequency"].fillna(merged["ssb_cu"])
        else:
            merged["frequency"] = None
        merged["source"] = fname
        merged["pci"] = merged["pci"].map(clean_num)
        merged["nRTAC"] = merged["nRTAC"].map(clean_num)
        merged["ssbFrequency"] = merged["ssbFrequency"].map(clean_num)
        if "frequency" not in merged.columns:
            merged["frequency"] = None
        merged["frequency"] = merged["frequency"].map(clean_num)
        merged["CGI"] = [make_cgi(a, b, c) for a, b, c in zip(merged["pLMNId"], merged["gNBId"], merged["cellLocalId"])]
        print(
            "  extracted",
            len(merged),
            "pci",
            merged["pci"].notna().sum(),
            "tac",
            merged["nRTAC"].notna().sum(),
            "ssb",
            merged["ssbFrequency"].notna().sum(),
            flush=True,
        )
        print("  CGI", merged["CGI"].dropna().head(3).tolist(), flush=True)
        nr_parts.append(
            merged[["source", "CGI", "gNBId", "cellLocalId", "pci", "nRTAC", "ssbFrequency", "frequency", "pLMNId", "gNBIdLength"]]
        )

    nr = pd.concat(nr_parts, ignore_index=True).drop_duplicates("CGI")
    print("5G去重", len(nr), "pci", nr["pci"].notna().sum(), "tac", nr["nRTAC"].notna().sum(), flush=True)

    print("\n解析 sdr.xlsx ...", flush=True)
    p = BASE / "sdr.xlsx"
    enb = read_cols(p, "ENBFunctionFDD", ["ManagedElement", "moId", "eNBId", "refPlmn"])
    cell = read_cols(
        p,
        "EUtranCellFDD",
        ["ManagedElement", "ldn", "moId", "cellLocalId", "refPlmn", "pci", "tac", "earfcnDl", "freqBandInd"],
    )
    print(" ENB/Cell", len(enb), len(cell), flush=True)

    def parent_enb(ldn):
        if ldn is None:
            return None
        m = re.search(r"ENBFunctionFDD=([^,]+)", str(ldn))
        return m.group(1) if m else None

    cell = cell.copy()
    cell["enbMo"] = cell["ldn"].map(parent_enb)
    enb2 = enb.rename(columns={"moId": "enbMo"})
    lte = cell.merge(enb2[["ManagedElement", "enbMo", "eNBId", "refPlmn"]], on=["ManagedElement", "enbMo"], how="left", suffixes=("_c", "_e"))
    if "refPlmn_c" in lte.columns:
        lte["plmn"] = lte["refPlmn_c"]
        if "refPlmn_e" in lte.columns:
            lte["plmn"] = lte["plmn"].fillna(lte["refPlmn_e"])
    elif "refPlmn" in lte.columns:
        lte["plmn"] = lte["refPlmn"]
    else:
        lte["plmn"] = None
    lte["pci"] = lte["pci"].map(clean_num)
    lte["tac"] = lte["tac"].map(clean_num)
    lte["earfcnDl"] = lte["earfcnDl"].map(clean_num)
    lte["CGI"] = [make_cgi(a, b, c) for a, b, c in zip(lte["plmn"], lte["eNBId"], lte["cellLocalId"])]
    print(
        "4G",
        len(lte),
        "pci",
        lte["pci"].notna().sum(),
        "tac",
        lte["tac"].notna().sum(),
        "earfcn",
        lte["earfcnDl"].notna().sum(),
        flush=True,
    )
    print("CGI", lte["CGI"].dropna().head(5).tolist(), "plmn", lte["plmn"].dropna().head(3).tolist(), flush=True)

    phys_keys = {}
    phys_nc = {}
    for cgi in phys["CGI"].astype(str):
        for v in variants(cgi):
            phys_keys[v] = cgi
        parts = re.split(r"[-_]", str(cgi).strip())
        if len(parts) >= 2 and parts[-2].isdigit() and parts[-1].isdigit():
            phys_nc[f"{int(parts[-2])}-{int(parts[-1])}"] = cgi

    def match_df(df, label):
        hit = 0
        hit_nc = 0
        for cgi in df["CGI"]:
            vs = variants(cgi)
            if any(v in phys_keys for v in vs):
                hit += 1
            else:
                parts = re.split(r"-", str(cgi or ""))
                if len(parts) == 4 and parts[2].isdigit() and parts[3].isdigit():
                    k = f"{int(parts[2])}-{int(parts[3])}"
                    if k in phys_nc:
                        hit_nc += 1
        print(
            f"{label}: n={len(df)} cgi非空={df['CGI'].notna().sum()} 精确命中={hit} node-cell命中={hit_nc}",
            flush=True,
        )

    print("\n匹配现网:", flush=True)
    match_df(nr, "5G")
    match_df(lte, "4G")
    for rat, sub in phys.groupby("网络制式"):
        print(f" 现网[{rat}] {len(sub)} 例", sub["CGI"].head(2).tolist(), flush=True)

    nr_map = {}
    lte_map = {}
    for _, r in nr.iterrows():
        for v in variants(r["CGI"]):
            nr_map[v] = r
        parts = re.split(r"-", str(r["CGI"] or ""))
        if len(parts) == 4 and parts[2].isdigit() and parts[3].isdigit():
            nr_map[f"NC:{int(parts[2])}-{int(parts[3])}"] = r
    for _, r in lte.iterrows():
        for v in variants(r["CGI"]):
            lte_map[v] = r
        parts = re.split(r"-", str(r["CGI"] or ""))
        if len(parts) == 4 and parts[2].isdigit() and parts[3].isdigit():
            lte_map[f"NC:{int(parts[2])}-{int(parts[3])}"] = r

    rows = []
    for _, r in phys.iterrows():
        cgi = str(r["CGI"])
        hit = None
        how = None
        for v in variants(cgi):
            if v in nr_map:
                hit = ("5G", nr_map[v])
                how = "cgi"
                break
            if v in lte_map:
                hit = ("4G", lte_map[v])
                how = "cgi"
                break
        if not hit:
            parts = re.split(r"[-_]", cgi.strip())
            if len(parts) >= 2 and parts[-2].isdigit() and parts[-1].isdigit():
                k = f"NC:{int(parts[-2])}-{int(parts[-1])}"
                if k in nr_map:
                    hit = ("5G", nr_map[k])
                    how = "nc"
                elif k in lte_map:
                    hit = ("4G", lte_map[k])
                    how = "nc"
        if not hit:
            continue
        kind, cfg = hit
        rows.append(
            {
                "CGI": cgi,
                "制式": r["网络制式"],
                "匹配": how,
                "现网频点": r["频点"],
                "网管PCI": cfg.get("pci"),
                "网管TAC": cfg.get("nRTAC") if kind == "5G" else cfg.get("tac"),
                "网管频点": (
                    cfg.get("ssbFrequency") if pd.notna(cfg.get("ssbFrequency")) else cfg.get("frequency")
                )
                if kind == "5G"
                else cfg.get("earfcnDl"),
                "来源": kind,
            }
        )

    cmp = pd.DataFrame(rows)
    print("\n命中现网行数", len(cmp), flush=True)
    if len(cmp):
        print("cgi匹配", (cmp["匹配"] == "cgi").sum(), "nc匹配", (cmp["匹配"] == "nc").sum(), flush=True)
        print(cmp.head(12).to_string(index=False), flush=True)
        print(
            "PCI非空",
            cmp["网管PCI"].notna().sum(),
            "TAC非空",
            cmp["网管TAC"].notna().sum(),
            "频点非空",
            cmp["网管频点"].notna().sum(),
            flush=True,
        )
        both = cmp.dropna(subset=["现网频点", "网管频点"]).copy()
        if len(both):
            both["现网频点"] = pd.to_numeric(both["现网频点"], errors="coerce")
            both["网管频点"] = pd.to_numeric(both["网管频点"], errors="coerce")
            eq = (both["现网频点"] == both["网管频点"]).sum()
            print(f"频点可比 {len(both)} 完全一致 {eq} ({eq/len(both)*100:.1f}%)", flush=True)
            mm = both[both["现网频点"] != both["网管频点"]].head(5)
            if len(mm):
                print(
                    "频点不一致样例:\n",
                    mm[["CGI", "制式", "现网频点", "网管频点", "网管PCI", "网管TAC"]].to_string(index=False),
                    flush=True,
                )
    else:
        print("无命中", flush=True)

    print("\n字段来源:", flush=True)
    print("5G: CellDefiningSSB.pci/ssbFrequency + NRCellDU.nRTAC", flush=True)
    print("4G: EUtranCellFDD.pci/tac/earfcnDl", flush=True)
    print("项目库: 原始小区表/物理表汇总 仅有频点，无PCI/TAC；无网管导入入口", flush=True)


if __name__ == "__main__":
    main()
