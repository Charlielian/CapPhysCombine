"""DuckDB / Excel IO utilities with Parquet intermediate cache.

缓存策略：
- 首次导入 Excel 时，解析结果写入 Parquet 缓存文件（.cache/parquet/）。
- 后续导入时，若 Excel 文件的 mtime+size 未变，直接从 Parquet 加载到 DuckDB，
  跳过 Excel 解析（提速 5-10 倍）。
- 缓存键：文件名 + mtime + size 的 hash。
- pyarrow 不可用时自动降级为直接解析 Excel（不影响功能）。

Excel 读取引擎：
- 默认使用 calamine（Rust 实现，比 openpyxl 快 5-10 倍）。
- calamine 不可用时自动降级为 openpyxl。
- 写入仍使用 openpyxl（calamine 仅支持读取）。
"""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path

import duckdb
import pandas as pd

from app.pipelines.paths import DB_PATH, PARQUET_CACHE_DIR, UNIFIED_DB_PATH

logger = logging.getLogger(__name__)

try:
    import pyarrow  # noqa: F401

    _HAS_PYARROW = True
except ImportError:
    _HAS_PYARROW = False

# ---------- Excel 读取引擎探测 ----------

_EXCEL_ENGINE: str = "openpyxl"  # 默认引擎


def _detect_excel_engine() -> str:
    """探测可用的最快 Excel 读取引擎。

    优先级：calamine > openpyxl。
    """
    try:
        import python_calamine  # noqa: F401

        return "calamine"
    except ImportError:
        pass
    return "openpyxl"


_EXCEL_ENGINE = _detect_excel_engine()


def get_excel_engine() -> str:
    """返回当前使用的 Excel 读取引擎名称。"""
    return _EXCEL_ENGINE


def get_db_connection() -> duckdb.DuckDBPyConnection:
    return duckdb.connect(str(DB_PATH))


def get_unified_db_connection() -> duckdb.DuckDBPyConnection:
    return duckdb.connect(str(UNIFIED_DB_PATH))


def init_db() -> None:
    if DB_PATH.exists():
        try:
            DB_PATH.unlink()
        except OSError:
            pass


def read_excel(path: Path, sheet_name: str | int | None = 0, engine: str | None = None) -> pd.DataFrame | dict[str, pd.DataFrame]:
    """读取 Excel 文件，优先使用 calamine（Rust 实现，快 5-10 倍）。

    Parameters
    ----------
    path : Path
        Excel 文件路径。
    sheet_name : str | int | None
        工作表名称或索引，默认读第一个。传 None 读全部。
    engine : str | None
        强制指定引擎。None 时自动探测（calamine > openpyxl）。
    """
    eng = engine or _EXCEL_ENGINE

    if sheet_name is None:
        # 读全部 sheet — calamine 通过 pd.ExcelFile 可以处理
        try:
            return pd.read_excel(path, engine=eng, sheet_name=None, dtype_backend="numpy_nullable")
        except Exception:
            if eng != "openpyxl":
                return pd.read_excel(path, engine="openpyxl", sheet_name=None, dtype_backend="numpy_nullable")
            raise

    try:
        return pd.read_excel(path, engine=eng, sheet_name=sheet_name, dtype_backend="numpy_nullable")
    except Exception:
        if eng != "openpyxl":
            # calamine 失败时降级到 openpyxl
            return pd.read_excel(path, engine="openpyxl", sheet_name=sheet_name, dtype_backend="numpy_nullable")
        raise


def _read_excel_flat(path: Path, engine: str | None = None) -> pd.DataFrame:
    """读取 Excel 第一个 sheet 为 DataFrame，始终返回单个 DataFrame。

    内部辅助函数，用于 engine 探测等不需要 sheet_name 参数的场景。
    """
    eng = engine or _EXCEL_ENGINE
    try:
        return pd.read_excel(path, engine=eng, dtype_backend="numpy_nullable")
    except Exception:
        if eng != "openpyxl":
            return pd.read_excel(path, engine="openpyxl", dtype_backend="numpy_nullable")
        raise


# ==============================================================================
# Parquet 缓存辅助
# ==============================================================================


def _cache_key(path: Path) -> str:
    """根据文件名 + mtime + size 生成缓存键。"""
    stat = path.stat()
    raw = f"{path.name}|{stat.st_mtime_ns}|{stat.st_size}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()[:16]


def _parquet_cache_path(path: Path) -> Path:
    """返回 Excel 文件对应的 Parquet 缓存路径。"""
    key = _cache_key(path)
    safe_stem = path.stem.replace("/", "_").replace("\\", "_")[:50]
    return PARQUET_CACHE_DIR / f"{safe_stem}_{key}.parquet"


def _is_cache_fresh(excel_path: Path, cache_path: Path) -> bool:
    """缓存文件存在且非空即视为可用（键已含 mtime+size，无需再校验）。"""
    if not _HAS_PYARROW:
        return False
    return cache_path.is_file() and cache_path.stat().st_size > 0


def _write_parquet_cache(df: pd.DataFrame, cache_path: Path) -> None:
    """将 DataFrame 写入 Parquet 缓存。"""
    if not _HAS_PYARROW:
        return
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(cache_path, index=False, engine="pyarrow")


def _load_parquet_to_table(
    conn: duckdb.DuckDBPyConnection,
    cache_path: Path,
    table_name: str,
    append: bool,
) -> int:
    """从 Parquet 文件直接加载到 DuckDB 表。"""
    parquet_uri = str(cache_path).replace("'", "''")
    if append:
        conn.execute(f'INSERT INTO "{table_name}" SELECT * FROM read_parquet(\'{parquet_uri}\')')
    else:
        conn.execute(f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM read_parquet(\'{parquet_uri}\')')
    # 取行数
    cnt = conn.execute(f'SELECT COUNT(*) FROM "{table_name}"').fetchone()
    return cnt[0] if cnt else 0


def clear_parquet_cache() -> int:
    """清空 Parquet 缓存目录，返回删除的文件数。"""
    if not PARQUET_CACHE_DIR.exists():
        return 0
    deleted = 0
    for f in PARQUET_CACHE_DIR.glob("*.parquet"):
        try:
            f.unlink()
            deleted += 1
        except OSError:
            pass
    return deleted


def cache_stats() -> dict:
    """返回缓存目录的统计信息。"""
    if not PARQUET_CACHE_DIR.exists():
        return {"exists": False, "file_count": 0, "total_size_bytes": 0}
    files = list(PARQUET_CACHE_DIR.glob("*.parquet"))
    total_size = sum(f.stat().st_size for f in files if f.is_file())
    return {
        "exists": True,
        "file_count": len(files),
        "total_size_bytes": total_size,
        "dir": str(PARQUET_CACHE_DIR),
    }


# ==============================================================================
# Excel -> DuckDB（带 Parquet 缓存）
# ==============================================================================


def excel_to_db(
    path: Path,
    table_name: str,
    conn: duckdb.DuckDBPyConnection,
    logger=None,
    chunk_size: int = 5000,
    append: bool = False,
) -> int:
    """大表分批导入：优先用 Parquet 缓存，未命中则解析 Excel 并写缓存。"""
    from app.pipelines.logging_util import GuiLogger

    logger = logger or GuiLogger()
    mode = "追加" if append else "导入"

    cache_path = _parquet_cache_path(path)
    if _is_cache_fresh(path, cache_path):
        logger.log(f"  [缓存] {mode} {path.name} -> {table_name} (Parquet 命中)")
        return _load_parquet_to_table(conn, cache_path, table_name, append)

    logger.log(f"  分批{mode} {path.name} -> {table_name} (每批 {chunk_size} 行, 引擎: {_EXCEL_ENGINE})")

    # 优先用 calamine 引擎（Rust 实现，快 5-10 倍）；降级到 openpyxl
    if _EXCEL_ENGINE == "calamine":
        try:
            from python_calamine import CalamineWorkbook

            wb = CalamineWorkbook.from_path(str(path))
            sheet_name_iter = iter(wb.sheet_names)
            first_sheet = next(sheet_name_iter, None)
            if first_sheet is None:
                return 0
            ws = wb.get_sheet_by_name(first_sheet)
            rows_iter = ws.iter_rows(values_only=True)
        except Exception:
            # calamine 失败，降级到 openpyxl
            from openpyxl import load_workbook

            wb = load_workbook(filename=str(path), read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
    else:
        from openpyxl import load_workbook

        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)

    try:
        headers = next(rows_iter)
    except StopIteration:
        if hasattr(wb, "close"):
            wb.close()
        return 0

    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(headers)]
    total_rows = 0
    chunk: list[tuple] = []
    first_write = not append
    # 同时收集所有 chunk 用于写 Parquet 缓存
    cache_chunks: list[pd.DataFrame] = []

    for row in rows_iter:
        chunk.append(row)
        if len(chunk) >= chunk_size:
            df_chunk = pd.DataFrame(chunk, columns=headers)
            cache_chunks.append(df_chunk)
            tmp_name = f"_tmp_{table_name}"
            conn.register(tmp_name, df_chunk)
            if first_write:
                conn.execute(f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM {tmp_name}')
                first_write = False
            else:
                conn.execute(f'INSERT INTO "{table_name}" SELECT * FROM {tmp_name}')
            total_rows += len(chunk)
            chunk = []

    if chunk:
        df_chunk = pd.DataFrame(chunk, columns=headers)
        cache_chunks.append(df_chunk)
        tmp_name = f"_tmp_{table_name}"
        conn.register(tmp_name, df_chunk)
        if first_write:
            conn.execute(f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM {tmp_name}')
        else:
            conn.execute(f'INSERT INTO "{table_name}" SELECT * FROM {tmp_name}')
        total_rows += len(chunk)

    wb.close()

    # 写 Parquet 缓存（合并所有 chunk）
    if cache_chunks:
        try:
            full_df = pd.concat(cache_chunks, ignore_index=True)
            _write_parquet_cache(full_df, cache_path)
            logger.log(f"  [缓存] 已写入 {cache_path.name}")
        except Exception as exc:
            logger.log(f"  [缓存] 写入失败（不影响导入）: {exc}")

    return total_rows


def small_excel_to_db(
    path: Path,
    table_name: str,
    conn: duckdb.DuckDBPyConnection,
    logger=None,
    append: bool = False,
) -> int:
    """小表导入：优先用 Parquet 缓存，未命中则解析 Excel 并写缓存。"""
    from app.pipelines.logging_util import GuiLogger

    logger = logger or GuiLogger()
    mode = "追加" if append else "导入"

    cache_path = _parquet_cache_path(path)
    if _is_cache_fresh(path, cache_path):
        logger.log(f"  [缓存] {mode} {path.name} -> {table_name} (Parquet 命中)")
        return _load_parquet_to_table(conn, cache_path, table_name, append)

    df = _read_excel_flat(path)
    df.columns = [str(c) for c in df.columns]

    tmp_name = f"_tmp_{table_name}"
    conn.register(tmp_name, df)
    if append:
        conn.execute(f'INSERT INTO "{table_name}" SELECT * FROM {tmp_name}')
    else:
        conn.execute(f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM {tmp_name}')

    logger.log(f"  {mode} {path.name} -> {table_name} ({len(df)} 行)")

    # 写 Parquet 缓存
    try:
        _write_parquet_cache(df, cache_path)
        logger.log(f"  [缓存] 已写入 {cache_path.name}")
    except Exception as exc:
        logger.log(f"  [缓存] 写入失败（不影响导入）: {exc}")

    return len(df)


def db_to_dataframe(query: str, conn: duckdb.DuckDBPyConnection) -> pd.DataFrame:
    return conn.execute(query).fetchdf()


def table_exists(conn: duckdb.DuckDBPyConnection, table_name: str) -> bool:
    result = conn.execute(
        f"SELECT 1 FROM information_schema.tables WHERE table_name = '{table_name}'"
    ).fetchone()
    return result is not None


def get_table_columns(conn: duckdb.DuckDBPyConnection, table_name: str) -> list[str]:
    result = conn.execute(f'DESCRIBE "{table_name}"').fetchdf()
    return list(result["column_name"])


def load_small_table(conn: duckdb.DuckDBPyConnection, name: str) -> pd.DataFrame:
    if not table_exists(conn, name):
        return pd.DataFrame()
    return db_to_dataframe(f'SELECT * FROM "{name}"', conn)


__all__ = [
    "get_db_connection",
    "get_unified_db_connection",
    "init_db",
    "read_excel",
    "get_excel_engine",
    "excel_to_db",
    "small_excel_to_db",
    "db_to_dataframe",
    "table_exists",
    "get_table_columns",
    "load_small_table",
    "clear_parquet_cache",
    "cache_stats",
]
